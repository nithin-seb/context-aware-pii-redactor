"""
Document protection module for PrivacyLens.

Provides three protection modes for sensitive PII in documents:
1. REDACT: Completely replaces sensitive values with category tokens (e.g. [REDACTED: EMAIL])
2. MASK: Partially masks values while preserving structural context (e.g. j***@example.com)
3. ANONYMIZE: Replaces sensitive values with consistent synthetic surrogate tokens (e.g. [PERSON_1])
"""

import io
import os
import re
from typing import Any, Dict, List, Optional, Set, Tuple, Union
import pymupdf
import docx

from .privacy_score import _extract_field

VALID_MODES = {"REDACT", "MASK", "ANONYMIZE"}

MIME_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain; charset=utf-8",
}


class ProtectionError(Exception):
    """Raised when document protection or sanitization fails."""
    pass


# Canonical category mapping for anonymization tokens and redact tags
CANONICAL_TYPES: Dict[str, str] = {
    "NAME": "PERSON",
    "PERSON": "PERSON",
    "PERSON_NAME": "PERSON",
    "INDIVIDUAL": "PERSON",
    "FULL_NAME": "PERSON",
    "FIRST_NAME": "PERSON",
    "LAST_NAME": "PERSON",
    "EMAIL": "EMAIL",
    "EMAIL_ADDRESS": "EMAIL",
    "PHONE": "PHONE",
    "PHONE_NUMBER": "PHONE",
    "TELEPHONE": "PHONE",
    "MOBILE": "PHONE",
    "CREDIT_CARD": "CREDIT_CARD",
    "CREDIT_CARD_NUMBER": "CREDIT_CARD",
    "CARD_NUMBER": "CREDIT_CARD",
    "DEBIT_CARD": "CREDIT_CARD",
    "SSN": "SSN",
    "SOCIAL_SECURITY_NUMBER": "SSN",
    "AADHAAR": "AADHAAR",
    "PAN": "PAN",
    "FINANCIAL_ID": "CREDIT_CARD",
    "CREDENTIAL": "PASSWORD",
    "PASSWORD": "PASSWORD",
    "API_KEY": "SECRET_KEY",
    "SECRET_KEY": "SECRET_KEY",
    "ADDRESS": "ADDRESS",
    "STREET_ADDRESS": "ADDRESS",
    "LOCATION": "ADDRESS",
    "PASSPORT": "PASSPORT",
    "BANK_ACCOUNT": "BANK_ACCOUNT",
    "DATE": "DATE",
    "DOB": "DATE",
    "DATE_OF_BIRTH": "DATE",
    "ORGANIZATION": "ORGANIZATION",
    "COMPANY": "ORGANIZATION",
    "IP_ADDRESS": "IP_ADDRESS",
    "MEDICAL": "MEDICAL_RECORD",
    "MEDICAL_RECORD": "MEDICAL_RECORD",
}


def _get_canonical_type(entity_type: Optional[str]) -> str:
    """Normalizes an entity type to a clean canonical uppercase label."""
    if not entity_type or not isinstance(entity_type, str):
        return "PII"
    cleaned = entity_type.strip().upper().replace(" ", "_").replace("-", "_")
    return CANONICAL_TYPES.get(cleaned, cleaned)


# ==============================================================================
# MASKING STRATEGIES
# ==============================================================================

def _mask_email(text: str) -> str:
    """
    Masks email while preserving initial character and domain:
    john@example.com -> j***@example.com
    j.smith@corp.org -> j***@corp.org
    """
    if "@" not in text:
        return _mask_generic(text)

    parts = text.split("@", 1)
    username, domain = parts[0], parts[1]

    if not username:
        return f"***@{domain}"

    first_char = username[0]
    return f"{first_char}***@{domain}"


def _mask_phone(text: str) -> str:
    """
    Masks phone numbers preserving formatting and the last 4 digits:
    9876543210 -> ******3210
    (555) 123-4567 -> (***) ***-4567
    """
    digits = [c for c in text if c.isdigit()]
    if len(digits) <= 4:
        kept = 1 if len(digits) > 1 else 0
    else:
        kept = 4

    digits_to_mask = len(digits) - kept
    masked_digit_count = 0
    result_chars = []

    for char in text:
        if char.isdigit():
            if masked_digit_count < digits_to_mask:
                result_chars.append("*")
                masked_digit_count += 1
            else:
                result_chars.append(char)
        else:
            result_chars.append(char)

    return "".join(result_chars)


def _mask_credit_card(text: str) -> str:
    """
    Masks credit card numbers preserving formatting and the last 4 digits:
    1234-5678-9012-3456 -> ****-****-****-3456
    1234567890123456 -> ************3456
    """
    return _mask_phone(text)


def _mask_ssn(text: str) -> str:
    """
    Masks SSN preserving the last 4 digits:
    123-45-6789 -> ***-**-6789
    123456789   -> *****6789
    """
    return _mask_phone(text)


def _mask_aadhaar(text: str) -> str:
    """
    Masks 12-digit Indian Aadhaar numbers:
    1234 5678 9012 -> **** **** 9012
    123456789012 -> ********9012
    """
    return _mask_phone(text)


def _mask_name(text: str) -> str:
    """
    Masks names by retaining initial letter per word:
    John Smith -> J*** S***
    Alice -> A***
    """
    words = text.split()
    if not words:
        return "***"

    masked_words = []
    for word in words:
        if len(word) == 0:
            continue
        first_letter = word[0]
        masked_words.append(f"{first_letter}***")

    return " ".join(masked_words)


def _mask_address(text: str) -> str:
    """
    Masks address street numbers and details:
    123 Main St -> *** Main St
    """
    masked = re.sub(r"^\d+", "***", text.strip())
    if masked == text.strip():
        masked = re.sub(r"\d+", "***", masked)
    return masked


def _mask_generic(text: str) -> str:
    """
    Safe generic masking fallback for arbitrary entity types:
    Keeps first and optionally last character with asterisks in between.
    """
    stripped = text.strip()
    length = len(stripped)

    if length <= 2:
        return "*" * max(1, length)
    elif length <= 4:
        return stripped[0] + "***"
    else:
        return stripped[0] + "***" + stripped[-1]


def get_masked_value(text: str, entity_type: Optional[str] = None) -> str:
    """Routes entity text to appropriate masking handler based on entity type."""
    canonical = _get_canonical_type(entity_type)

    if canonical == "EMAIL":
        return _mask_email(text)
    elif canonical in ("PHONE", "MOBILE"):
        return _mask_phone(text)
    elif canonical in ("CREDIT_CARD", "FINANCIAL_ID", "PAN"):
        return _mask_credit_card(text)
    elif canonical in ("SSN", "AADHAAR"):
        return _mask_ssn(text)
    elif canonical == "PERSON":
        return _mask_name(text)
    elif canonical == "ADDRESS":
        return _mask_address(text)
    else:
        return _mask_generic(text)


# ==============================================================================
# SPAN PREPARATION & RESOLUTION
# ==============================================================================

def _extract_and_resolve_spans(
    original_text: str,
    entities: List[Any],
) -> List[Dict[str, Any]]:
    """Extracts, validates, and sorts non-overlapping spans from detected entities."""
    if not original_text or not entities:
        return []

    doc_len = len(original_text)
    candidate_spans: List[Dict[str, Any]] = []

    for entity in entities:
        start_val = _extract_field(entity, "start", "start_offset", "start_pos", "begin")
        end_val = _extract_field(entity, "end", "end_offset", "end_pos")
        text_val = _extract_field(entity, "text", "value", "content", "raw")
        raw_type = _extract_field(entity, "entity_type", "type", "label", "category") or "PII"

        if start_val is None and end_val is None and not text_val:
            continue

        start_int: Optional[int] = None
        end_int: Optional[int] = None
        try:
            if start_val is not None:
                start_int = int(start_val)
            if end_val is not None:
                end_int = int(end_val)
        except (ValueError, TypeError):
            start_int = None
            end_int = None

        if start_int is None or end_int is None or start_int < 0 or end_int > doc_len or start_int >= end_int:
            if text_val and isinstance(text_val, str) and text_val in original_text:
                search_idx = 0
                while search_idx < doc_len:
                    found_idx = original_text.find(text_val, search_idx)
                    if found_idx == -1:
                        break
                    candidate_spans.append({
                        "start": found_idx,
                        "end": found_idx + len(text_val),
                        "text": text_val,
                        "entity_type": raw_type,
                        "original_entity": entity,
                    })
                    search_idx = found_idx + len(text_val)
            continue

        span_text = original_text[start_int:end_int]
        extracted_text = str(text_val) if text_val is not None else span_text

        candidate_spans.append({
            "start": start_int,
            "end": end_int,
            "text": extracted_text,
            "entity_type": raw_type,
            "original_entity": entity,
        })

    if not candidate_spans:
        return []

    sorted_candidates = sorted(
        candidate_spans,
        key=lambda s: (s["start"], -(s["end"] - s["start"])),
    )

    resolved: List[Dict[str, Any]] = []
    last_end = -1

    for span in sorted_candidates:
        if span["start"] >= last_end:
            resolved.append(span)
            last_end = span["end"]

    return resolved


class DocumentProtector:
    """
    Manages document protection state and executes safe replacements.
    Maintains anonymization mappings consistently across occurrences.
    """

    def __init__(self) -> None:
        self.type_counters: Dict[str, int] = {}
        self.anonymize_mapping: Dict[Tuple[str, str], str] = {}

    def get_anonymized_placeholder(self, text: str, entity_type: str) -> str:
        """Returns a consistent anonymized placeholder for a given entity value and type."""
        canonical = _get_canonical_type(entity_type)
        lookup_key = (canonical, text.strip().lower())

        if lookup_key in self.anonymize_mapping:
            return self.anonymize_mapping[lookup_key]

        current_count = self.type_counters.get(canonical, 0) + 1
        self.type_counters[canonical] = current_count

        placeholder = f"[{canonical}_{current_count}]"
        self.anonymize_mapping[lookup_key] = placeholder
        return placeholder

    def protect(
        self,
        original_text: str,
        entities: Optional[List[Any]] = None,
        mode: Union[str, Any] = "redact",
    ) -> Dict[str, Any]:
        """Protects document text using the specified mode: 'redact', 'mask', or 'anonymize'."""
        if not original_text:
            return {
                "protected_text": "",
                "mode": str(mode).lower(),
                "entities_protected": 0,
                "anonymization_mapping": {},
            }

        if not entities:
            return {
                "protected_text": original_text,
                "mode": str(mode).lower(),
                "entities_protected": 0,
                "anonymization_mapping": {},
            }

        mode_str = str(mode).strip().lower()
        if hasattr(mode, "value"):
            mode_str = str(mode.value).strip().lower()

        resolved_spans = _extract_and_resolve_spans(original_text, entities)
        if not resolved_spans:
            return {
                "protected_text": original_text,
                "mode": mode_str,
                "entities_protected": 0,
                "anonymization_mapping": {},
            }

        replacement_actions: List[Dict[str, Any]] = []
        for span in resolved_spans:
            entity_text = span["text"]
            entity_type = span["entity_type"]

            if mode_str == "mask":
                replacement = get_masked_value(entity_text, entity_type)
            elif mode_str == "anonymize":
                replacement = self.get_anonymized_placeholder(entity_text, entity_type)
            else:
                canonical = _get_canonical_type(entity_type)
                if canonical == "PII":
                    replacement = "[REDACTED]"
                else:
                    replacement = f"[REDACTED: {canonical}]"

            replacement_actions.append({
                "start": span["start"],
                "end": span["end"],
                "replacement": replacement,
            })

        protected_text = original_text
        for action in sorted(replacement_actions, key=lambda a: a["start"], reverse=True):
            start = action["start"]
            end = action["end"]
            rep = action["replacement"]
            protected_text = protected_text[:start] + rep + protected_text[end:]

        exportable_mapping = {
            k[1]: v for k, v in self.anonymize_mapping.items()
        }

        return {
            "protected_text": protected_text,
            "mode": mode_str,
            "entities_protected": len(resolved_spans),
            "anonymization_mapping": exportable_mapping,
        }


# ==============================================================================
# FILE-LEVEL (PDF, DOCX, TXT) PROTECTION FOR FASTAPI / DOWNLOADS
# ==============================================================================

def _build_replacement_map(
    entities: List[Dict[str, Any]],
    mode: str
) -> Dict[str, str]:
    """Build replacement dictionary of entity value -> protected token for file sanitization."""
    replacements: Dict[str, str] = {}
    protector = DocumentProtector()
    mode_lower = mode.strip().lower()

    sorted_entities = sorted(
        entities,
        key=lambda x: len(str(_extract_field(x, "value", "text", default=""))),
        reverse=True,
    )

    for entity in sorted_entities:
        orig_val = str(_extract_field(entity, "value", "text", default="")).strip()
        if not orig_val or orig_val in replacements:
            continue

        entity_type = str(_extract_field(entity, "type", "entity_type", default="PII")).upper()

        if mode_lower == "redact":
            canonical = _get_canonical_type(entity_type)
            replacement = f"[REDACTED: {canonical}]"
        elif mode_lower == "mask":
            replacement = get_masked_value(orig_val, entity_type)
        elif mode_lower == "anonymize":
            canonical = _get_canonical_type(entity_type)
            if canonical == "PERSON":
                replacement = protector.get_anonymized_placeholder(orig_val, "PERSON")
            elif canonical == "EMAIL":
                counter = protector.type_counters.get("EMAIL", 0) + 1
                protector.type_counters["EMAIL"] = counter
                replacement = f"user{counter}@privacylens.internal"
            elif canonical == "PHONE":
                counter = protector.type_counters.get("PHONE", 0) + 1
                protector.type_counters["PHONE"] = counter
                replacement = f"+1-555-01{counter:02d}"
            else:
                replacement = protector.get_anonymized_placeholder(orig_val, canonical)
        else:
            replacement = f"[PROTECTED_{entity_type}]"

        replacements[orig_val] = replacement

    return replacements


def protect_txt(raw_bytes: bytes, replacement_map: Dict[str, str]) -> bytes:
    """Sanitize TXT document."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            text = raw_bytes.decode(encoding)
            for orig_val, rep in replacement_map.items():
                if orig_val in text:
                    text = text.replace(orig_val, rep)
            return text.encode("utf-8")
        except UnicodeDecodeError:
            continue
    raise ProtectionError("Failed to decode text document for protection.")


def protect_docx(raw_bytes: bytes, replacement_map: Dict[str, str]) -> bytes:
    """Sanitize DOCX document preserving format."""
    try:
        doc = docx.Document(io.BytesIO(raw_bytes))

        for p in doc.paragraphs:
            for orig_val, rep in replacement_map.items():
                if orig_val in p.text:
                    p.text = p.text.replace(orig_val, rep)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for orig_val, rep in replacement_map.items():
                        if orig_val in cell.text:
                            cell.text = cell.text.replace(orig_val, rep)

        output_stream = io.BytesIO()
        doc.save(output_stream)
        return output_stream.getvalue()
    except Exception as e:
        raise ProtectionError(f"Failed to protect DOCX document: {str(e)}")


def protect_pdf(
    raw_bytes: bytes,
    replacement_map: Dict[str, str],
    mode: str
) -> bytes:
    """Sanitize PDF document using PyMuPDF native redactions."""
    try:
        doc = pymupdf.open(stream=raw_bytes, filetype="pdf")

        for page in doc:
            for orig_val, rep in replacement_map.items():
                if not orig_val:
                    continue
                rects = page.search_for(orig_val)
                for rect in rects:
                    if mode.upper() == "REDACT":
                        page.add_redact_annot(
                            rect,
                            text=rep,
                            fontsize=9,
                            fill=(0, 0, 0),
                            text_color=(1, 1, 1),
                        )
                    else:
                        page.add_redact_annot(
                            rect,
                            text=rep,
                            fontsize=9,
                            fill=(0.95, 0.95, 0.95),
                            text_color=(0.1, 0.1, 0.1),
                        )

            page.apply_redactions()

        output_bytes = doc.tobytes(deflate=True, clean=True)
        doc.close()
        return output_bytes
    except Exception as e:
        if 'doc' in locals() and not doc.is_closed:
            doc.close()
        raise ProtectionError(f"Failed to protect PDF document: {str(e)}")


# ==============================================================================
# UNIFIED PROTECT_DOCUMENT DISPATCHER
# ==============================================================================

def protect_document(
    target: Optional[Union[str, Dict[str, Any]]] = None,
    entities: Optional[List[Any]] = None,
    mode: str = "redact",
    scan_data: Optional[Dict[str, Any]] = None,
) -> Union[str, Tuple[bytes, str, str]]:
    """
    Main protection interface supporting both string text protection and file-level scan protection.

    Usage A (String text):
        protected_str = protect_document("Email: john@example.com", entities=[...], mode="redact")

    Usage B (File scan dict):
        protected_bytes, filename, mime_type = protect_document(scan_data=scan_data_dict, mode="REDACT")
    """
    # Usage B: scan_data dictionary passed
    actual_scan_data = scan_data if scan_data is not None else (target if isinstance(target, dict) else None)
    if actual_scan_data is not None:
        scan_data_obj = actual_scan_data
        mode_upper = mode.upper().strip()
        if mode_upper not in VALID_MODES:
            raise ProtectionError(f"Invalid protection mode '{mode}'. Must be REDACT, MASK, or ANONYMIZE.")

        scan_entities = scan_data_obj.get("entities", [])
        file_type = scan_data_obj.get("file_type", "txt").lower()
        filename = scan_data_obj.get("filename", f"document.{file_type}")
        raw_bytes = scan_data_obj.get("raw_bytes")
        doc_text = scan_data_obj.get("text", "")

        if not raw_bytes and doc_text:
            raw_bytes = doc_text.encode("utf-8")
            file_type = "txt"

        if not raw_bytes:
            raise ProtectionError("No document data available for protection.")

        replacement_map = _build_replacement_map(scan_entities, mode_upper)

        base_name, ext = os.path.splitext(filename)
        output_filename = f"{base_name}_{mode_upper.lower()}{ext if ext else f'.{file_type}'}"
        mime_type = MIME_TYPES.get(file_type, "application/octet-stream")

        if file_type == "pdf":
            protected_bytes = protect_pdf(raw_bytes, replacement_map, mode_upper)
        elif file_type == "docx":
            protected_bytes = protect_docx(raw_bytes, replacement_map)
        elif file_type == "txt":
            protected_bytes = protect_txt(raw_bytes, replacement_map)
        else:
            raise ProtectionError(f"Unsupported file type for protection: {file_type}")

        return protected_bytes, output_filename, mime_type

    # Usage A: string text passed
    original_text = str(target)
    protector = DocumentProtector()
    result = protector.protect(original_text, entities, mode=mode)
    return result["protected_text"]


def redact_text(
    original_text: str,
    entities: Optional[List[Any]] = None,
) -> str:
    """Convenience wrapper for REDACT mode."""
    return str(protect_document(original_text, entities, mode="redact"))


def mask_text(
    original_text: str,
    entities: Optional[List[Any]] = None,
) -> str:
    """Convenience wrapper for MASK mode."""
    return str(protect_document(original_text, entities, mode="mask"))


def anonymize_text(
    original_text: str,
    entities: Optional[List[Any]] = None,
) -> str:
    """Convenience wrapper for ANONYMIZE mode."""
    return str(protect_document(original_text, entities, mode="anonymize"))
