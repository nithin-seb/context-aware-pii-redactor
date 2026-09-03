"""Document parser service for extracting plain text from PDF, DOCX, and TXT files."""

import io
import os
from typing import Any, Dict, Optional
import pymupdf
import docx

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB

SUPPORTED_EXTENSIONS = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".txt": "txt",
}

SUPPORTED_CONTENT_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "txt",
}


class DocumentParserError(Exception):
    """Base exception for document parsing errors."""
    pass


class FileTooLargeError(DocumentParserError):
    """Raised when uploaded file exceeds the maximum size limit."""
    pass


class UnsupportedFileTypeError(DocumentParserError):
    """Raised when uploaded file format is not supported."""
    pass


class EmptyDocumentError(DocumentParserError):
    """Raised when uploaded document has no extractable text."""
    pass


class ExtractionError(DocumentParserError):
    """Raised when document content extraction fails."""
    pass


def detect_file_type(filename: str, content_type: Optional[str] = None) -> str:
    """Detect file type from extension and optional MIME content type."""
    ext = os.path.splitext(filename or "")[1].lower()
    if ext in SUPPORTED_EXTENSIONS:
        return SUPPORTED_EXTENSIONS[ext]

    if content_type:
        normalized_mime = content_type.split(";")[0].strip().lower()
        if normalized_mime in SUPPORTED_CONTENT_TYPES:
            return SUPPORTED_CONTENT_TYPES[normalized_mime]

    raise UnsupportedFileTypeError(
        f"Unsupported file type. Supported types are PDF (.pdf), Word (.docx), and Plain Text (.txt)."
    )


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF using PyMuPDF without OCR."""
    try:
        doc = pymupdf.open(stream=file_bytes, filetype="pdf")
    except Exception as e:
        raise ExtractionError(f"Failed to open PDF document: {str(e)}")

    try:
        pages_text = []
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text")
            if text:
                pages_text.append(text.strip())
        doc.close()
        return "\n\n".join(pages_text)
    except Exception as e:
        if not doc.is_closed:
            doc.close()
        raise ExtractionError(f"Failed to extract text from PDF: {str(e)}")


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX paragraphs and tables using python-docx."""
    try:
        doc_stream = io.BytesIO(file_bytes)
        doc = docx.Document(doc_stream)
    except Exception as e:
        raise ExtractionError(f"Failed to parse DOCX document: {str(e)}")

    try:
        extracted_chunks = []

        # Extract text from paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                extracted_chunks.append(text)

        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    extracted_chunks.append(" | ".join(row_cells))

        return "\n".join(extracted_chunks)
    except Exception as e:
        raise ExtractionError(f"Failed to extract text from DOCX: {str(e)}")


def extract_text_from_txt(file_bytes: bytes) -> str:
    """Extract text from TXT with multiple encoding fallbacks."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            return file_bytes.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    raise ExtractionError("Unable to decode text file with standard encodings.")


def extract_document(
    file_bytes: bytes,
    filename: str,
    content_type: Optional[str] = None
) -> Dict[str, Any]:
    """
    Extract clean text and metadata from an uploaded document.

    Args:
        file_bytes: Raw bytes of the uploaded file.
        filename: Name of the uploaded file.
        content_type: Optional MIME content type.

    Returns:
        Dict with keys: text, filename, file_type, size_bytes

    Raises:
        FileTooLargeError
        UnsupportedFileTypeError
        EmptyDocumentError
        ExtractionError
    """
    if not file_bytes or len(file_bytes) == 0:
        raise EmptyDocumentError("The uploaded file is empty (0 bytes).")

    size_bytes = len(file_bytes)
    if size_bytes > MAX_FILE_SIZE_BYTES:
        max_mb = MAX_FILE_SIZE_BYTES // (1024 * 1024)
        raise FileTooLargeError(
            f"File size ({size_bytes / (1024 * 1024):.2f} MB) exceeds maximum allowed limit of {max_mb} MB."
        )

    file_type = detect_file_type(filename, content_type)

    if file_type == "pdf":
        extracted = extract_text_from_pdf(file_bytes)
    elif file_type == "docx":
        extracted = extract_text_from_docx(file_bytes)
    elif file_type == "txt":
        extracted = extract_text_from_txt(file_bytes)
    else:
        raise UnsupportedFileTypeError(f"Unsupported file type '{file_type}'.")

    # Clean text
    clean_text = extracted.replace("\r\n", "\n").replace("\r", "\n").strip()

    if not clean_text:
        raise EmptyDocumentError(
            "No extractable text found in the document. Image-only or empty files are not supported."
        )

    return {
        "text": clean_text,
        "filename": filename,
        "file_type": file_type,
        "size_bytes": size_bytes,
    }
